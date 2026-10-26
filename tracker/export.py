"""Export both trackers to one .xlsx workbook (openpyxl), with approved changes highlighted.

Sheets: "Budget", "Work plan", "Change log". A changed cell gets a yellow fill and
a comment naming the proposal, the old value and the approver.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings
from tracker import db

HEADER_FILL = PatternFill("solid", fgColor="1F3B4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CHANGED_FILL = PatternFill("solid", fgColor="FFE58A")

BUDGET_COLS = ["line_id", "program", "category", "annual_budget", "spent_to_date", "committed", "owner"]
WORKPLAN_COLS = ["id", "program", "activity", "owner", "due_date", "status", "percent_complete", "notes", "depends_on"]
MONEY = {"annual_budget", "spent_to_date", "committed"}


def _sheet(ws, cols, rows, key, tracker, log) -> None:
    ws.append([c.replace("_", " ").title() for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    latest = {}
    for c in log:
        if c["tracker"] == tracker:
            latest[(c["row_id"], c["field"])] = c  # later rows win
    for r_idx, row in enumerate(rows, start=2):
        ws.append([row[c] for c in cols])
        for c_idx, col in enumerate(cols, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if col in MONEY:
                cell.number_format = "#,##0.00"
            change = latest.get((row[key], col))
            if change:
                cell.fill = CHANGED_FILL
                cell.comment = Comment(f"Changed by {change['proposal_id']}: {change['old_value']} -> "
                                       f"{change['new_value']}. Approved by {change['approved_by']}.", "Good Measure")
    for i, col in enumerate(cols, start=1):
        width = max(len(col), *(len(str(r[col] or "")) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 48)
    ws.freeze_panes = "A2"


def build_workbook() -> bytes:
    log = db.change_log()
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    _sheet(ws, BUDGET_COLS, db.budget_lines(), "line_id", "budget", log)
    _sheet(wb.create_sheet("Work plan"), WORKPLAN_COLS, db.activities(), "id", "workplan", log)
    cl = wb.create_sheet("Change log")
    cl.append(["Tracker", "Row", "Field", "Old value", "New value", "Proposal", "Approved by", "Applied at"])
    for cell in cl[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for c in log:
        cl.append([c["tracker"], c["row_id"], c["field"], c["old_value"], c["new_value"], c["proposal_id"],
                   c["approved_by"], c["applied_at"]])
    cl.append([])
    cl.append([f"{settings.ORG_NAME} (fictional). Demo date {settings.DEMO_TODAY}. Yellow cells were changed "
               "by a human-approved proposal."])
    cl.cell(row=cl.max_row, column=1).alignment = Alignment(wrap_text=False)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
